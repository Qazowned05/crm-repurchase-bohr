from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.modules.auth.models import User
from app.modules.customers.models import Customer
from app.modules.imports.models import ImportJob, ImportRowError
from app.modules.products.models import Product, ProductRepurchaseRule
from app.modules.sales.models import Sale, SaleItem
from app.modules.imports.routes import excel_safe


def make_user(db: Session, email: str, role: str) -> User:
    record = User(email=email, full_name=email, password_hash=hash_password("correct-password"), role=role, is_active=True)
    db.add(record)
    db.commit()
    return record


def auth(client: TestClient, email: str) -> dict[str, str]:
    response = client.post("/api/v1/auth/login", json={"email": email, "password": "correct-password"})
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def workbook(headers: list[str], rows: list[list[object]], formula: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if formula:
        sheet["B2"] = "=1+1"
    output = BytesIO(); book.save(output)
    return output.getvalue()


CUSTOMER_HEADERS = ["dni", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district", "responsible_advisor_email"]
PRODUCT_HEADERS = ["code", "name", "brand", "category", "unit_price", "is_active", "duration_days", "alert_days", "effective_from"]
HISTORICAL_SALES_HEADERS = ["numero_venta", "customer_dni", "sale_date", "advisor_email", "channel", "channel_detail", "product_code", "quantity", "notes", "historical_unit_price"]


def preview(client: TestClient, headers: dict[str, str], import_type: str, content: bytes, **data: str) -> dict:
    response = client.post(f"/api/v1/imports/{import_type}/preview", headers=headers, data=data, files={"file": ("upload.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 201, response.text
    return response.json()


def test_import_templates_and_import_access_are_supervisor_only(client: TestClient, db: Session) -> None:
    advisor = make_user(db, "advisor-import@example.com", "ASESOR")
    denied = client.get("/api/v1/imports/templates/customers", headers=auth(client, advisor.email))
    assert denied.status_code == 403
    supervisor = make_user(db, "supervisor-import@example.com", "SUPERVISOR")
    template = client.get("/api/v1/imports/templates/customers", headers=auth(client, supervisor.email))
    assert template.status_code == 200
    assert template.content.startswith(b"PK")
    historical = client.get("/api/v1/imports/templates/historical_sales", headers=auth(client, supervisor.email))
    workbook_template = load_workbook(BytesIO(historical.content))
    assert workbook_template["Data"].column_dimensions["J"].hidden is True
    assert workbook_template["Data"].cell(1, 10).value == "historical_unit_price"


def test_preview_persists_formula_and_duplicate_errors(client: TestClient, db: Session) -> None:
    supervisor = make_user(db, "preview@example.com", "SUPERVISOR")
    result = preview(client, auth(client, supervisor.email), "customers", workbook(CUSTOMER_HEADERS, [["70000001", "Ana", "Torres", "999999999", "ana@example.com", "", "", "", "", "", ""], ["70000001", "Beto", "Diaz", "999999998", "beto@example.com", "", "", "", "", "", ""]], formula=True))
    assert result["state"] == "PREVIEW_ERRORS"
    assert result["rejected_rows"] >= 1
    assert db.query(ImportJob).count() == 1
    assert db.query(ImportRowError).count() >= 1
    assert client.post(f"/api/v1/imports/{result['id']}/commit", headers=auth(client, supervisor.email)).status_code == 409


def test_excel_exports_escape_formula_prefixes() -> None:
    assert excel_safe("=1+1") == "'=1+1"
    assert excel_safe("+cmd") == "'+cmd"
    assert excel_safe("normal text") == "normal text"


def test_clean_import_commits_customers_and_versioned_product_rule(client: TestClient, db: Session) -> None:
    supervisor = make_user(db, "commit@example.com", "SUPERVISOR")
    headers = auth(client, supervisor.email)
    customer_job = preview(client, headers, "customers", workbook(CUSTOMER_HEADERS, [["70000001", "Ana", "Torres", "999999999", "ana@example.com", "", "", "", "", "", ""]]))
    committed = client.post(f"/api/v1/imports/{customer_job['id']}/commit", headers=headers)
    assert committed.status_code == 200
    assert committed.json()["state"] == "COMMITTED"
    assert db.query(Customer).count() == 1
    assert client.post("/api/v1/brands", headers=headers, json={"name": "Bohr"}).status_code == 201
    assert client.post("/api/v1/product-categories", headers=headers, json={"name": "Suplementos"}).status_code == 201
    product_job = preview(client, headers, "products", workbook(PRODUCT_HEADERS, [["COL-001", "Colageno", "Bohr", "Suplementos", "19.90", "true", "30", "15,5", "2026-01-01"]]))
    assert client.post(f"/api/v1/imports/{product_job['id']}/commit", headers=headers).status_code == 200
    assert db.query(Product).count() == 1
    assert db.query(ProductRepurchaseRule).count() == 1


def test_product_import_rejects_unknown_catalog_values(client: TestClient, db: Session) -> None:
    supervisor = make_user(db, "catalog-import@example.com", "SUPERVISOR")
    result = preview(client, auth(client, supervisor.email), "products", workbook(PRODUCT_HEADERS, [["COL-404", "Colageno", "Desconocida", "Suplementos", "19.90", "true", "", "", ""]]))
    assert result["state"] == "PREVIEW_ERRORS"
    assert "existing active catalog" in result["errors"][0]["message"]


def test_historical_sales_import_groups_lines_and_preserves_price(client: TestClient, db: Session) -> None:
    supervisor = make_user(db, "historic-supervisor@example.com", "SUPERVISOR")
    advisor = make_user(db, "historic-advisor@example.com", "ASESOR")
    headers = auth(client, supervisor.email)
    customer_job = preview(client, headers, "customers", workbook(CUSTOMER_HEADERS, [["70000001", "Ana", "Torres", "999999999", "ana@example.com", "", "", "", "", "", advisor.email]]))
    assert client.post(f"/api/v1/imports/{customer_job['id']}/commit", headers=headers).status_code == 200
    assert client.post("/api/v1/brands", headers=headers, json={"name": "Bohr"}).status_code == 201
    assert client.post("/api/v1/product-categories", headers=headers, json={"name": "Suplementos"}).status_code == 201
    product_job = preview(client, headers, "products", workbook(PRODUCT_HEADERS, [["HIST-001", "Producto", "Bohr", "Suplementos", "99.90", "true", "30", "15,5", "2025-01-01"]]))
    assert client.post(f"/api/v1/imports/{product_job['id']}/commit", headers=headers).status_code == 200
    job = preview(client, headers, "historical_sales", workbook(HISTORICAL_SALES_HEADERS, [["V-0001", "70000001", "2026-01-01", advisor.email, "DIGITAL", "", "HIST-001", "2", "migrada", "42.50"]]))
    assert job["state"] == "PREVIEW_READY"
    assert client.post(f"/api/v1/imports/{job['id']}/commit", headers=headers).status_code == 200
    assert db.query(Sale).count() == 1
    item = db.query(SaleItem).one()
    assert item.quantity == 2
    assert str(item.unit_price) == "42.50"
    assert db.query(Customer).one().acquisition_channel == "DIGITAL"
    retry = preview(client, headers, "historical_sales", workbook(HISTORICAL_SALES_HEADERS, [["V-0001", "70000001", "2026-01-01", advisor.email, "DIGITAL", "", "HIST-001", "2", "migrada", "42.50"]]))
    assert retry["state"] == "PREVIEW_READY"
    assert retry["new_rows"] == 0
    assert client.post(f"/api/v1/imports/{retry['id']}/commit", headers=headers).status_code == 200
    assert db.query(Sale).count() == 1
    assert db.query(Sale).one().external_reference == "V-0001"
