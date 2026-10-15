from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font
from pydantic import ValidationError
from sqlalchemy import select, tuple_
from sqlalchemy.orm import Session

from app.core.audit import record_audit
from app.core.config import settings
from app.core.database import get_db
from app.dependencies import require_roles
from app.modules.auth.models import User
from app.modules.customers.models import Customer
from app.modules.customers.schemas import CustomerCreate
from app.modules.customers.routes import transfer_customer_assignment
from app.modules.imports.models import ImportJob, ImportRowError
from app.modules.products.models import Brand, Product, ProductCategory, ProductRepurchaseRule
from app.modules.products.schemas import RuleCreate
from app.modules.sales.models import Sale, SaleItem
from app.modules.sales.routes import recompute_chain, rule_for
from app.modules.alerts.service import run_alert_generation
from app.data.ubigeo import locations, location_hierarchy

router = APIRouter(prefix="/api/v1/imports", tags=["imports"])

HEADERS = {
    "customers": ["customer_type", "dni", "business_name", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district", "responsible_advisor_email"],
    "products": ["code", "name", "brand", "category", "unit_price", "is_active", "duration_days", "alert_days", "effective_from"],
    "historical_sales": ["numero_venta", "customer_dni", "sale_date", "advisor_email", "channel", "channel_detail", "product_code", "quantity", "notes", "historical_unit_price"],
}
LEGACY_CUSTOMER_HEADERS = ["dni", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district", "responsible_advisor_email"]


def job_response(job: ImportJob, db: Session) -> dict:
    errors = list(db.scalars(select(ImportRowError).where(ImportRowError.import_job_id == job.id).order_by(ImportRowError.row_number)))
    return {"id": job.id, "type": job.import_type, "state": job.state, "total_rows": job.total_rows,
            "valid_rows": job.valid_rows, "new_rows": job.new_rows, "update_rows": job.update_rows,
            "rejected_rows": job.rejected_rows, "errors": [{"row_number": e.row_number, "field": e.field, "message": e.message} for e in errors]}


def add_error(errors: list[dict], row: int | None, message: str, field: str | None = None) -> None:
    errors.append({"row_number": row, "field": field, "message": message})


def cell_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_date(value: str) -> date:
    """Accept ISO dates, Excel timestamps, and DD/MM/YYYY spreadsheet dates."""
    try:
        return date.fromisoformat(value)
    except ValueError:
        try:
            return datetime.fromisoformat(value).date()
        except ValueError:
            return datetime.strptime(value, "%d/%m/%Y").date()


def excel_safe(value: object) -> object:
    """Prevent Excel from evaluating user-controlled text as a formula."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def catalog_record(model: type[Brand] | type[ProductCategory], name: str, label: str, db: Session) -> Brand | ProductCategory:
    record = db.scalar(select(model).where(model.name == name.strip()))
    if record is None or not record.is_active:
        raise ValueError(f"{label} must be an existing active catalog entry")
    return record


def workbook_rows(content: bytes, import_type: str) -> tuple[list[dict], list[dict]]:
    errors: list[dict] = []
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            if any(name.lower().endswith(("vbaproject.bin", ".bin")) for name in names):
                return [], [{"row_number": None, "field": None, "message": "Macro-enabled workbooks are not allowed"}]
            if sum(info.file_size for info in archive.infolist()) > settings.import_max_file_size * 10:
                return [], [{"row_number": None, "field": None, "message": "Workbook expands beyond the allowed limit"}]
    except BadZipFile:
        return [], [{"row_number": None, "field": None, "message": "Invalid or corrupt XLSX file"}]
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False, keep_vba=False)
    except Exception:
        return [], [{"row_number": None, "field": None, "message": "Invalid or corrupt XLSX file"}]
    if workbook.security and workbook.security.lockStructure:
        return [], [{"row_number": None, "field": None, "message": "Protected workbooks are not allowed"}]
    if "Data" not in workbook.sheetnames:
        return [], [{"row_number": None, "field": None, "message": "Workbook must contain a Data sheet"}]
    sheet = workbook["Data"]
    if sheet.protection.sheet:
        return [], [{"row_number": None, "field": None, "message": "Protected worksheets are not allowed"}]
    headers = [cell_text(cell.value) for cell in sheet[1]]
    legacy_customers = import_type == "customers" and headers == LEGACY_CUSTOMER_HEADERS
    if headers != HEADERS[import_type] and not legacy_customers:
        return [], [{"row_number": 1, "field": None, "message": "Headers must exactly match the template"}]
    rows: list[dict] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if all(cell.value is None for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            add_error(errors, row_number, "Formulas are not allowed in data fields")
            continue
        if len(rows) >= settings.import_max_rows:
            add_error(errors, row_number, f"Workbook exceeds {settings.import_max_rows} data rows")
            break
        row = {header: cell_text(cell.value) for header, cell in zip(headers, cells, strict=False)}
        if legacy_customers:
            row.update({"customer_type": "PERSONA", "business_name": None})
        rows.append({"row_number": row_number, "data": row})
    return rows, errors


def validate_rows(rows: list[dict], import_type: str, db: Session, allow_updates: bool, allow_reassignment: bool) -> tuple[list[dict], list[dict]]:
    if import_type == "historical_sales":
        return validate_historical_sales(rows, db)
    errors: list[dict] = []
    seen: set[str] = set()
    valid: list[dict] = []
    for source in rows:
        row_number, data = source["row_number"], source["data"]
        key = data["dni"] if import_type == "customers" else data["code"]
        key = key.upper() if key else None
        if not key:
            add_error(errors, row_number, "Identifier is required", "dni" if import_type == "customers" else "code")
            continue
        if key in seen:
            add_error(errors, row_number, "Duplicate identifier within workbook", "dni" if import_type == "customers" else "code")
            continue
        seen.add(key)
        try:
            if import_type == "customers":
                data["dni"] = key
                existing = db.scalar(select(Customer).where(Customer.dni == key))
                action = "update" if existing else "new"
                if existing and not allow_updates:
                    raise ValueError("Customer already exists; enable allow_updates to update it")
                if not existing:
                    customer_payload = CustomerCreate(**{field: data[field] for field in ("dni", "customer_type", "business_name", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district")})
                else:
                    customer_payload = CustomerCreate(**{
                        "dni": data["dni"],
                        "customer_type": data["customer_type"] or existing.customer_type,
                        "business_name": data["business_name"] if data["business_name"] is not None else existing.business_name,
                        "first_names": data["first_names"] or existing.first_names,
                        "last_names": data["last_names"] or existing.last_names,
                        "phone": data["phone"] or existing.phone,
                        "email": data["email"] if data["email"] is not None else existing.email,
                        "birth_date": data["birth_date"] if data["birth_date"] is not None else existing.birth_date,
                        "condition": data["condition"] if data["condition"] is not None else existing.condition,
                        "department": data["department"] if data["department"] is not None else existing.department,
                        "province": data["province"] if data["province"] is not None else existing.province,
                        "district": data["district"] if data["district"] is not None else existing.district,
                    })
                data.update(customer_payload.model_dump(mode="json"))
                location_fields = ("department", "province", "district")
                supplied_location = [data[field] is not None for field in location_fields]
                if any(supplied_location) and not all(supplied_location):
                    raise ValueError("department, province, and district must be provided together")
                if all(supplied_location):
                    location = next((row for row in locations() if all(data[field] == row[f"{field}_name"] for field in location_fields)), None)
                    if location is None:
                        raise ValueError("department, province, and district must match the official INEI catalog")
                    data["ubigeo"] = location["ubigeo"]
                advisor_email = data.pop("responsible_advisor_email", None)
                if advisor_email:
                    advisor = db.scalar(select(User).where(User.email == advisor_email.lower()))
                    if advisor is None or not advisor.is_active or advisor.role != "ASESOR":
                        raise ValueError("Responsible advisor must be an active ASESOR")
                    if existing and advisor.id != existing.responsible_advisor_id and not allow_reassignment:
                        raise ValueError("Enable allow_reassignment to change the responsible advisor")
                    data["responsible_advisor_id"] = advisor.id
            else:
                data["code"] = key
                existing = db.scalar(select(Product).where(Product.code == key))
                action = "update" if existing else "new"
                if existing and not allow_updates:
                    raise ValueError("Product code already exists; enable allow_updates to update it")
                if not existing:
                    if not data["name"] or not data["brand"] or not data["category"] or not data["unit_price"]:
                        raise ValueError("name, brand, category, and unit_price are required for new products")
                    catalog_record(Brand, data["brand"], "Brand", db)
                    catalog_record(ProductCategory, data["category"], "Category", db)
                else:
                    if data["brand"]:
                        catalog_record(Brand, data["brand"], "Brand", db)
                    if data["category"]:
                        catalog_record(ProductCategory, data["category"], "Category", db)
                data["is_active"] = parse_bool(data["is_active"]) if data["is_active"] else None
                if data["unit_price"] is not None:
                    data["unit_price"] = str(Decimal(data["unit_price"]))
                rule_fields = ["duration_days", "alert_days", "effective_from"]
                supplied = [field for field in rule_fields if data[field] is not None]
                if supplied and len(supplied) != len(rule_fields):
                    raise ValueError("All repurchase-rule fields are required when importing a rule")
                if supplied:
                    data["duration_days"] = int(data["duration_days"])
                    data["alert_days"] = [int(day.strip()) for day in data["alert_days"].split(",")]
                    effective_from = import_date(data["effective_from"])
                    data["effective_from"] = effective_from.isoformat()
                    RuleCreate(**{field: data[field] for field in rule_fields})
                    if existing and db.scalar(select(ProductRepurchaseRule).where(ProductRepurchaseRule.product_id == existing.id, ProductRepurchaseRule.effective_from == effective_from)) is not None:
                        raise ValueError("A repurchase rule already exists for this effective date")
            valid.append({"row_number": row_number, "action": action, "data": data})
        except (ValidationError, ValueError, TypeError) as error:
            add_error(errors, row_number, str(error))
    return valid, errors


def validate_historical_sales(rows: list[dict], db: Session) -> tuple[list[dict], list[dict]]:
    errors: list[dict] = []
    valid: list[dict] = []
    sales: dict[str, dict] = {}
    customer_dnis = {(row["data"]["customer_dni"] or "").upper() for row in rows}
    advisor_emails = {(row["data"]["advisor_email"] or "").lower() for row in rows}
    product_codes = {(row["data"]["product_code"] or "").upper() for row in rows}
    references = {(row["data"]["numero_venta"] or "").upper() for row in rows}
    customers = {
        customer.dni: customer
        for customer in db.scalars(select(Customer).where(Customer.dni.in_(customer_dnis)))
    }
    advisors = {
        advisor.email: advisor
        for advisor in db.scalars(select(User).where(User.email.in_(advisor_emails)))
    }
    products = {
        product.code: product
        for product in db.scalars(select(Product).where(Product.code.in_(product_codes)))
    }
    rules_by_product: dict[str, list[ProductRepurchaseRule]] = {}
    for rule in db.scalars(
        select(ProductRepurchaseRule)
        .where(ProductRepurchaseRule.product_id.in_([product.id for product in products.values()]))
        .order_by(ProductRepurchaseRule.product_id, ProductRepurchaseRule.effective_from.desc())
    ):
        rules_by_product.setdefault(rule.product_id, []).append(rule)
    existing_sales = {
        sale.external_reference: sale
        for sale in db.scalars(select(Sale).where(Sale.external_reference.in_(references)))
    }
    existing_items = {
        (item.sale_id, item.product_id): item
        for item in db.scalars(
            select(SaleItem)
            .join(Sale, Sale.id == SaleItem.sale_id)
            .where(Sale.external_reference.in_(references))
        )
    }
    for source in rows:
        row_number, data = source["row_number"], source["data"]
        try:
            reference = (data["numero_venta"] or "").upper()
            if not reference:
                raise ValueError("numero_venta is required")
            customer = customers.get((data["customer_dni"] or "").upper())
            advisor = advisors.get((data["advisor_email"] or "").lower())
            product = products.get((data["product_code"] or "").upper())
            sale_date = import_date(data["sale_date"] or "")
            quantity = int(data["quantity"] or "")
            channel = (data["channel"] or "").upper()
            if customer is None or customer.status != "ACTIVO": raise ValueError("customer_dni must identify an active customer")
            if advisor is None or advisor.role != "ASESOR" or not advisor.is_active: raise ValueError("advisor_email must identify an active ASESOR")
            if product is None or not product.is_active: raise ValueError("product_code must identify an active product")
            if sale_date >= date.today(): raise ValueError("sale_date must be earlier than today for a historical import")
            if quantity <= 0: raise ValueError("quantity must be positive")
            historical_unit_price = Decimal(data["historical_unit_price"]) if data["historical_unit_price"] else product.unit_price
            if historical_unit_price is None or historical_unit_price <= 0: raise ValueError("historical_unit_price must be positive when provided")
            if channel not in {"TV", "DIGITAL", "OTROS"}: raise ValueError("channel must be TV, DIGITAL, or OTROS")
            if channel == "OTROS" and not data["channel_detail"]: raise ValueError("channel_detail is required for OTROS")
            if channel != "OTROS" and data["channel_detail"]: raise ValueError("channel_detail is only allowed for OTROS")
            if not any(rule.effective_from <= sale_date for rule in rules_by_product.get(product.id, [])):
                raise ValueError("product has no repurchase rule effective on sale_date")
            sale = sales.setdefault(reference, {"customer_id": customer.id, "advisor_id": advisor.id, "sale_date": sale_date.isoformat(), "channel": channel, "channel_detail": data["channel_detail"], "notes": data["notes"], "products": set()})
            if any(sale[field] != value for field, value in (("customer_id", customer.id), ("advisor_id", advisor.id), ("sale_date", sale_date.isoformat()), ("channel", channel), ("channel_detail", data["channel_detail"]), ("notes", data["notes"]))):
                raise ValueError("all rows with the same numero_venta must share customer, date, advisor, channel, detail, and notes")
            if product.id in sale["products"]: raise ValueError("a product may only appear once per numero_venta")
            sale["products"].add(product.id)
            existing_sale = existing_sales.get(reference)
            if existing_sale is not None:
                fields = (
                    ("customer_id", customer.id), ("advisor_id", advisor.id), ("sale_date", sale_date),
                    ("acquisition_channel", channel), ("acquisition_channel_detail", data["channel_detail"]), ("notes", data["notes"]),
                )
                if any(getattr(existing_sale, field) != value for field, value in fields):
                    raise ValueError("numero_venta already belongs to a different sale")
                existing_item = existing_items.get((existing_sale.id, product.id))
                if existing_item is None or existing_item.quantity != quantity or existing_item.unit_price != historical_unit_price:
                    raise ValueError("numero_venta already exists with different products, quantities, or prices")
                action = "existing"
            else:
                action = "new"
            valid.append({"row_number": row_number, "action": action, "data": {"numero_venta": reference, "customer_id": customer.id, "advisor_id": advisor.id, "sale_date": sale_date.isoformat(), "channel": channel, "channel_detail": data["channel_detail"], "notes": data["notes"], "product_id": product.id, "quantity": quantity, "unit_price": str(historical_unit_price)}})
        except (ValueError, TypeError, ArithmeticError) as error:
            add_error(errors, row_number, str(error))
    return valid, errors


def parse_bool(value: str) -> bool:
    if value.lower() in {"true", "1", "yes", "si", "sí"}:
        return True
    if value.lower() in {"false", "0", "no"}:
        return False
    raise ValueError("is_active must be true or false")


@router.get("/templates/{import_type}")
def download_template(import_type: str, _: User = Depends(require_roles("SUPERVISOR", "ADMIN")), db: Session = Depends(get_db)) -> StreamingResponse:
    if import_type not in HEADERS:
        raise HTTPException(status_code=404, detail="Unknown import type")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(HEADERS[import_type])
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    lists = workbook.create_sheet("_lists")
    lists.sheet_state = "hidden"
    example = workbook.create_sheet("Ejemplo")
    example.append(HEADERS[import_type])
    for cell in example[1]:
        cell.font = Font(bold=True)
    if import_type == "products":
        brands = list(db.scalars(select(Brand).where(Brand.is_active.is_(True)).order_by(Brand.name)))
        categories = list(db.scalars(select(ProductCategory).where(ProductCategory.is_active.is_(True)).order_by(ProductCategory.name)))
        for row, brand in enumerate(brands, start=1): lists.cell(row=row, column=1, value=excel_safe(brand.name))
        for row, category in enumerate(categories, start=1): lists.cell(row=row, column=2, value=excel_safe(category.name))
        if brands:
            validation = DataValidation(type="list", formula1=f"'_lists'!$A$1:$A${len(brands)}", allow_blank=False)
            sheet.add_data_validation(validation); validation.add("C2:C1000")
        if categories:
            validation = DataValidation(type="list", formula1=f"'_lists'!$B$1:$B${len(categories)}", allow_blank=False)
            sheet.add_data_validation(validation); validation.add("D2:D1000")
        example.append([
            "EJEMPLO-001", "Producto de ejemplo", excel_safe(brands[0].name) if brands else "Marca existente",
            excel_safe(categories[0].name) if categories else "Categoría existente", "19.90", "true", "30", "15,5", date.today().isoformat(),
        ])
    elif import_type == "customers":
        advisors = list(db.scalars(select(User).where(User.role == "ASESOR", User.is_active.is_(True)).order_by(User.email)))
        for row, advisor in enumerate(advisors, start=1): lists.cell(row=row, column=1, value=excel_safe(advisor.email))
        if advisors:
            validation = DataValidation(type="list", formula1=f"'_lists'!$A$1:$A${len(advisors)}", allow_blank=False)
            sheet.add_data_validation(validation); validation.add("M2:M1000")
        example.append([
            "PERSONA", "70000001", "", "Ana", "Ejemplo", "999999999", "ana@example.com", "1985-06-15", "",
            "LIMA", "LIMA", "MIRAFLORES", advisors[0].email if advisors else "asesor@example.com",
        ])
        # Excel validations use named ranges because dependent dropdowns cannot directly
        # reference another worksheet. The helper sheet remains hidden from operators.
        hierarchy = location_hierarchy()
        for row, department in enumerate(hierarchy, start=1):
            lists.cell(row=row, column=3, value=department["name"])
            lists.cell(row=row, column=4, value=department["code"])
        workbook.defined_names.add(DefinedName("departments", attr_text=f"'_lists'!$C$1:$C${len(hierarchy)}"))
        workbook.defined_names.add(DefinedName("department_codes", attr_text=f"'_lists'!$C$1:$D${len(hierarchy)}"))
        next_column = 5
        for department in hierarchy:
            province_column = get_column_letter(next_column)
            for row, province in enumerate(department["provinces"], start=1):
                lists.cell(row=row, column=next_column, value=province["name"])
            workbook.defined_names.add(DefinedName(
                f"dep_{department['code']}",
                attr_text=f"'_lists'!${province_column}$1:${province_column}${len(department['provinces'])}",
            ))
            next_column += 1
            province_codes_column = get_column_letter(next_column)
            province_codes_end = next_column + 1
            for row, province in enumerate(department["provinces"], start=1):
                lists.cell(row=row, column=next_column, value=province["name"])
                lists.cell(row=row, column=province_codes_end, value=province["code"])
            workbook.defined_names.add(DefinedName(
                f"province_codes_{department['code']}",
                attr_text=f"'_lists'!${province_codes_column}$1:${get_column_letter(province_codes_end)}${len(department['provinces'])}",
            ))
            next_column += 2
            for province in department["provinces"]:
                district_column = get_column_letter(next_column)
                for row, district in enumerate(province["districts"], start=1):
                    lists.cell(row=row, column=next_column, value=district["name"])
                workbook.defined_names.add(DefinedName(
                    f"prov_{province['code']}",
                    attr_text=f"'_lists'!${district_column}$1:${district_column}${len(province['districts'])}",
                ))
                next_column += 1
        department_validation = DataValidation(type="list", formula1="=departments", allow_blank=False)
        customer_type_validation = DataValidation(type="list", formula1='"PERSONA,EMPRESA"', allow_blank=False)
        province_validation = DataValidation(type="list", formula1='=INDIRECT("dep_"&VLOOKUP(J2,department_codes,2,FALSE))', allow_blank=False)
        district_validation = DataValidation(type="list", formula1='=INDIRECT("prov_"&VLOOKUP(K2,INDIRECT("province_codes_"&VLOOKUP(J2,department_codes,2,FALSE)),2,FALSE))', allow_blank=False)
        sheet.add_data_validation(customer_type_validation); customer_type_validation.add("A2:A1000")
        sheet.add_data_validation(department_validation); department_validation.add("J2:J1000")
        sheet.add_data_validation(province_validation); province_validation.add("K2:K1000")
        sheet.add_data_validation(district_validation); district_validation.add("L2:L1000")
    else:
        advisors = list(db.scalars(select(User).where(User.role == "ASESOR", User.is_active.is_(True)).order_by(User.email)))
        products = list(db.scalars(select(Product).where(Product.is_active.is_(True)).order_by(Product.code)))
        for row, advisor in enumerate(advisors, start=1): lists.cell(row=row, column=1, value=excel_safe(advisor.email))
        for row, product in enumerate(products, start=1): lists.cell(row=row, column=2, value=excel_safe(product.code))
        if advisors:
            validation = DataValidation(type="list", formula1=f"'_lists'!$A$1:$A${len(advisors)}", allow_blank=False)
            sheet.add_data_validation(validation); validation.add("D2:D1000")
        if products:
            validation = DataValidation(type="list", formula1=f"'_lists'!$B$1:$B${len(products)}", allow_blank=False)
            sheet.add_data_validation(validation); validation.add("G2:G1000")
        example.append(["V-0001", "70000001", (date.today() - timedelta(days=1)).isoformat(), excel_safe(advisors[0].email) if advisors else "asesor@example.com", "DIGITAL", "", excel_safe(products[0].code) if products else "CODIGO-PRODUCTO", "1", "Venta histórica migrada", ""])
        sheet.column_dimensions["J"].hidden = True
        example.column_dimensions["J"].hidden = True
        instructions = workbook.create_sheet("Instrucciones")
        instructions.append(["Campo", "Cómo usarlo"])
        instructions.append(["numero_venta", "Número o código de la venta original. Repítelo solo si esa misma venta tiene más de un producto."])
        instructions.append(["product_code", "Selecciona un producto disponible desde la lista desplegable."])
        instructions.append(["historical_unit_price", "Columna técnica oculta. Una integración puede incluir el precio histórico original; si está vacía, se usa el precio vigente del catálogo al validar."])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{import_type}_template.xlsx"'})


@router.post("/{import_type}/preview", status_code=status.HTTP_201_CREATED)
async def preview_import(import_type: str, file: UploadFile = File(...), allow_updates: bool = Form(False), allow_reassignment: bool = Form(False), current_user: User = Depends(require_roles("SUPERVISOR", "ADMIN")), db: Session = Depends(get_db)) -> dict:
    if import_type not in HEADERS:
        raise HTTPException(status_code=404, detail="Unknown import type")
    filename = file.filename or "upload.xlsx"
    content = await file.read(settings.import_max_file_size + 1)
    errors: list[dict] = []
    rows: list[dict] = []
    if not filename.lower().endswith(".xlsx"):
        add_error(errors, None, "Only .xlsx files are accepted")
    elif len(content) > settings.import_max_file_size:
        add_error(errors, None, f"File exceeds {settings.import_max_file_size} byte limit")
    else:
        rows, errors = workbook_rows(content, import_type)
        if not errors:
            valid, validation_errors = validate_rows(rows, import_type, db, allow_updates, allow_reassignment)
            errors.extend(validation_errors)
        else:
            valid = []
    job = ImportJob(import_type=import_type, created_by_user_id=current_user.id, source_filename=filename, source_file=content, allow_updates=allow_updates, allow_reassignment=allow_reassignment, total_rows=len(rows), valid_rows=len(valid), new_rows=sum(row["action"] == "new" for row in valid), update_rows=sum(row["action"] == "update" for row in valid), rejected_rows=len(errors), rows_data=valid, state="PREVIEW_ERRORS" if errors else "PREVIEW_READY")
    db.add(job)
    db.flush()
    for error in errors:
        db.add(ImportRowError(import_job_id=job.id, **error))
    record_audit(db, actor_id=current_user.id, entity_type="import_job", entity_id=job.id, action="PREVIEWED", after={"type": import_type, "state": job.state, "total_rows": job.total_rows})
    db.commit()
    return job_response(job, db)


@router.get("/{job_id}")
def get_import_job(job_id: str, current_user: User = Depends(require_roles("SUPERVISOR", "ADMIN")), db: Session = Depends(get_db)) -> dict:
    job = db.get(ImportJob, job_id)
    if job is None or (job.created_by_user_id != current_user.id and current_user.role != "ADMIN"):
        raise HTTPException(status_code=404, detail="Import job not found")
    return job_response(job, db)


@router.get("/{job_id}/errors")
def download_errors(job_id: str, current_user: User = Depends(require_roles("SUPERVISOR", "ADMIN")), db: Session = Depends(get_db)) -> StreamingResponse:
    job = db.get(ImportJob, job_id)
    if job is None or (job.created_by_user_id != current_user.id and current_user.role != "ADMIN"):
        raise HTTPException(status_code=404, detail="Import job not found")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Errors"
    sheet.append(["row_number", "field", "message"])
    for error in db.scalars(select(ImportRowError).where(ImportRowError.import_job_id == job.id).order_by(ImportRowError.row_number)):
        sheet.append([excel_safe(error.row_number), excel_safe(error.field), excel_safe(error.message)])
    output = BytesIO(); workbook.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="import_{job.id}_errors.xlsx"'})


def commit_historical_sales(rows: list[dict], current_user: User, job: ImportJob, db: Session) -> None:
    grouped_sales: dict[str, list[dict]] = {}
    for row in rows:
        grouped_sales.setdefault(row["data"]["numero_venta"], []).append(row["data"])

    references = set(grouped_sales)
    existing_references = set(db.scalars(select(Sale.external_reference).where(Sale.external_reference.in_(references))))
    grouped_sales = {reference: lines for reference, lines in grouped_sales.items() if reference not in existing_references}
    product_ids = {line["product_id"] for lines in grouped_sales.values() for line in lines}
    rules_by_product: dict[str, list[ProductRepurchaseRule]] = {}
    for rule in db.scalars(
        select(ProductRepurchaseRule)
        .where(ProductRepurchaseRule.product_id.in_(product_ids))
        .order_by(ProductRepurchaseRule.product_id, ProductRepurchaseRule.effective_from.desc())
    ):
        rules_by_product.setdefault(rule.product_id, []).append(rule)

    sales_by_reference: dict[str, Sale] = {}
    for reference, lines in grouped_sales.items():
        data = lines[0]
        sale = Sale(
            customer_id=data["customer_id"], advisor_id=data["advisor_id"], sale_date=date.fromisoformat(data["sale_date"]),
            notes=data["notes"], status="CONFIRMADA", acquisition_channel=data["channel"],
            acquisition_channel_detail=data["channel_detail"], external_reference=reference,
        )
        sales_by_reference[reference] = sale
    db.add_all(sales_by_reference.values())
    db.flush()

    items: list[SaleItem] = []
    affected_pairs: set[tuple[str, str]] = set()
    for reference, lines in grouped_sales.items():
        sale = sales_by_reference[reference]
        for line in lines:
            rule = next((rule for rule in rules_by_product[line["product_id"]] if rule.effective_from <= sale.sale_date), None)
            if rule is None:
                raise ValueError(f"Product {line['product_id']} no longer has an effective repurchase rule")
            items.append(SaleItem(
                sale_id=sale.id, product_id=line["product_id"], quantity=line["quantity"],
                unit_price=Decimal(line["unit_price"]), rule_duration_days=rule.duration_days,
                rule_alert_days=rule.alert_days, expected_repurchase_date=sale.sale_date + timedelta(days=rule.duration_days),
            ))
            affected_pairs.add((sale.customer_id, line["product_id"]))
        record_audit(db, actor_id=current_user.id, entity_type="sale", entity_id=sale.id, action="HISTORICAL_IMPORT", after={"numero_venta": reference, "import_job_id": job.id})
    db.add_all(items)
    db.flush()

    chains: dict[tuple[str, str], list[SaleItem]] = {}
    pair_list = list(affected_pairs)
    for start in range(0, len(pair_list), 300):
        for item, sale in db.execute(
            select(SaleItem, Sale)
            .join(Sale, Sale.id == SaleItem.sale_id)
            .where(Sale.status == "CONFIRMADA", tuple_(Sale.customer_id, SaleItem.product_id).in_(pair_list[start:start + 300]))
            .order_by(Sale.customer_id, SaleItem.product_id, Sale.sale_date, Sale.created_at, SaleItem.created_at)
        ):
            chains.setdefault((sale.customer_id, item.product_id), []).append(item)
    for chain in chains.values():
        previous: SaleItem | None = None
        for item in chain:
            item.purchase_type = "COMPRA" if previous is None else "RECOMPRA"
            item.prior_confirmed_item_id = previous.id if previous else None
            previous = item

    imported_sales = list(sales_by_reference.values())
    customers = {
        customer.id: customer
        for customer in db.scalars(select(Customer).where(Customer.id.in_({sale.customer_id for sale in imported_sales})))
    }
    for sale in sorted(imported_sales, key=lambda current: (current.sale_date, current.created_at)):
        customer = customers[sale.customer_id]
        if customer.acquisition_channel is None:
            customer.acquisition_channel = sale.acquisition_channel
            customer.acquisition_channel_detail = sale.acquisition_channel_detail
    run_alert_generation(date.today(), db, commit=False)


@router.post("/{job_id}/commit")
def commit_import(job_id: str, current_user: User = Depends(require_roles("SUPERVISOR", "ADMIN")), db: Session = Depends(get_db)) -> dict:
    job = db.get(ImportJob, job_id)
    if job is None or (job.created_by_user_id != current_user.id and current_user.role != "ADMIN"):
        raise HTTPException(status_code=404, detail="Import job not found")
    if job.state != "PREVIEW_READY":
        raise HTTPException(status_code=409, detail="Only a clean preview can be committed")
    try:
        if job.import_type == "historical_sales":
            commit_historical_sales(job.rows_data, current_user, job, db)
        processed_references: set[str] = set()
        for row in job.rows_data:
            data = row["data"]
            if job.import_type == "historical_sales":
                continue
            if job.import_type == "customers":
                customer = db.scalar(select(Customer).where(Customer.dni == data["dni"]))
                if customer is None:
                    customer = Customer(**{field: data[field] for field in ("dni", "customer_type", "business_name", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district", "ubigeo", "responsible_advisor_id") if data.get(field) is not None})
                    db.add(customer); db.flush()
                    record_audit(db, actor_id=current_user.id, entity_type="customer", entity_id=customer.id, action="CREATED", after={"dni": customer.dni, "import_job_id": job.id})
                elif job.allow_updates:
                    before = {}; changes = {field: data[field] for field in ("customer_type", "business_name", "first_names", "last_names", "phone", "email", "birth_date", "condition", "department", "province", "district", "ubigeo") if data.get(field) is not None}
                    if "email" in changes:
                        changes["email"] = changes["email"].lower()
                    for field, value in changes.items(): before[field] = getattr(customer, field); setattr(customer, field, value)
                    if changes: record_audit(db, actor_id=current_user.id, entity_type="customer", entity_id=customer.id, action="UPDATED", before=before, after=changes)
                    if data.get("responsible_advisor_id") is not None:
                        transfer_customer_assignment(customer, data["responsible_advisor_id"], "Reasignación mediante importación", current_user.id, db)
                else: raise ValueError(f"Customer {data['dni']} now conflicts with an existing record")
            elif job.import_type == "products":
                product = db.scalar(select(Product).where(Product.code == data["code"]))
                if product is None:
                    brand = catalog_record(Brand, data["brand"], "Brand", db)
                    category = catalog_record(ProductCategory, data["category"], "Category", db)
                    product = Product(code=data["code"], name=data["name"].strip(), brand_id=brand.id, category_id=category.id, unit_price=data["unit_price"], is_active=True if data["is_active"] is None else data["is_active"])
                    db.add(product); db.flush(); record_audit(db, actor_id=current_user.id, entity_type="product", entity_id=product.id, action="CREATED", after={"code": product.code, "import_job_id": job.id})
                elif job.allow_updates:
                    changes = {field: data[field] for field in ("name", "unit_price", "is_active") if data.get(field) is not None}
                    if data.get("brand") is not None:
                        changes["brand_id"] = catalog_record(Brand, data["brand"], "Brand", db).id
                    if data.get("category") is not None:
                        changes["category_id"] = catalog_record(ProductCategory, data["category"], "Category", db).id
                    before = {field: getattr(product, field) for field in changes}
                    for field, value in changes.items(): setattr(product, field, value)
                    if changes: record_audit(db, actor_id=current_user.id, entity_type="product", entity_id=product.id, action="UPDATED", before=before, after=changes)
                else: raise ValueError(f"Product {data['code']} now conflicts with an existing record")
                if data.get("duration_days") is not None:
                    rule = ProductRepurchaseRule(product_id=product.id, duration_days=data["duration_days"], alert_days=data["alert_days"], effective_from=date.fromisoformat(data["effective_from"]), created_by_user_id=current_user.id)
                    db.add(rule); db.flush(); record_audit(db, actor_id=current_user.id, entity_type="product_rule", entity_id=rule.id, action="CREATED", after={"product_id": product.id, "import_job_id": job.id})
            else:
                # Rows are validated individually but committed as grouped sale lines.
                if data["numero_venta"] in processed_references:
                    continue
                grouped = [row["data"] for row in job.rows_data if row["data"]["numero_venta"] == data["numero_venta"]]
                processed_references.add(data["numero_venta"])
                if db.scalar(select(Sale.id).where(Sale.external_reference == data["numero_venta"])) is not None:
                    continue
                sale = Sale(customer_id=data["customer_id"], advisor_id=data["advisor_id"], sale_date=date.fromisoformat(data["sale_date"]), notes=data["notes"], status="CONFIRMADA", acquisition_channel=data["channel"], acquisition_channel_detail=data["channel_detail"], external_reference=data["numero_venta"])
                db.add(sale); db.flush()
                product_ids = set()
                for line in grouped:
                    rule = rule_for(line["product_id"], sale.sale_date, db)
                    db.add(SaleItem(sale_id=sale.id, product_id=line["product_id"], quantity=line["quantity"], unit_price=Decimal(line["unit_price"]), rule_duration_days=rule.duration_days, rule_alert_days=rule.alert_days, expected_repurchase_date=sale.sale_date + timedelta(days=rule.duration_days)))
                    product_ids.add(line["product_id"])
                db.flush(); recompute_chain(sale.customer_id, product_ids, db)
                record_audit(db, actor_id=current_user.id, entity_type="sale", entity_id=sale.id, action="HISTORICAL_IMPORT", after={"numero_venta": data["numero_venta"], "import_job_id": job.id})
        job.state = "COMMITTED"; job.committed_at = datetime.now(timezone.utc)
        record_audit(db, actor_id=current_user.id, entity_type="import_job", entity_id=job.id, action="COMMITTED", after={"total_rows": job.total_rows})
        db.commit()
    except Exception as error:
        db.rollback()
        raise HTTPException(status_code=409, detail=f"Import could not be committed atomically: {error}") from error
    return job_response(job, db)
